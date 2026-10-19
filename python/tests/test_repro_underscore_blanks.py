"""
Form-regeneration regressions (vela-wood patch, 2026-07-30 / 2026-08-14).

1. Fill-in blanks ($[_______], signature lines) were eaten by the inline
   italic parser as empty `_..._` pairs.
2. The word diff split `**_Term_**` across hunks; each unbalanced half lost
   its formatting.
3. (a) accept-all counted pPr tab-stop DEFINITIONS as content and kept
   fully-deleted paragraphs as empty shells; (b) positional (diff-pinned)
   whole-paragraph heading replacements skipped the PARAGRAPH_REPLACE path.
4. Heading/list content lost trailing spaces, failing post-apply verification.
"""

import io
import re

import docx
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Inches

from adeu.diff import _words_to_chars, generate_edits_from_text
from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine

BLANK = "$[_______]"
SIGNATURE_LINE = "By: ____________________"
WORD_SPLIT = r"(\s+|\w+|[^\w\s])"


def _engine(*paragraphs) -> RedlineEngine:
    # Each item is "text" or ("text", "Style Name").
    doc = docx.Document()
    for item in paragraphs:
        text, style = item if isinstance(item, tuple) else (item, None)
        doc.add_paragraph(text, style=style)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return RedlineEngine(stream, author="QA Bot")


def _accepted_paragraphs(engine: RedlineEngine):
    engine.accept_all_revisions()
    result = docx.Document(engine.save_to_stream())
    return [(p.style.name, p.text) for p in result.paragraphs]


def _tokens(text: str):
    _, _, token_array = _words_to_chars(text, text)
    return token_array


# --- Fix 1: blanks survive the inline parser -------------------------------


def test_inline_parser_keeps_fill_in_blanks():
    engine = _engine("x")

    for text in (f"Price: {BLANK} per share", SIGNATURE_LINE, "__"):
        runs = engine._parse_inline_markdown(text)

        assert "".join(t for t, _ in runs) == text
        assert all(not props for _, props in runs)


def test_inline_parser_still_parses_real_italics():
    engine = _engine("x")

    assert engine._parse_inline_markdown("the _Company_ shall") == [
        ("the ", {}),
        ("Company", {"italic": True}),
        (" shall", {}),
    ]
    assert engine._parse_inline_markdown("**_Company_**") == [
        ("Company", {"bold": True, "italic": True}),
    ]


def test_blank_inserted_via_edit_keeps_underscores():
    engine = _engine("Price: TBD per share.")
    edit = ModifyText(type="modify", target_text="TBD", new_text=BLANK)

    engine.process_batch([edit])

    assert _accepted_paragraphs(engine) == [("Normal", f"Price: {BLANK} per share.")]


# --- Fix 2: emphasis spans are atomic diff tokens ---------------------------


def test_emphasis_span_is_one_diff_token():
    assert "**_Company_**" in _tokens("the **_Company_** shall")
    assert "_Company_" in _tokens("the _Company_ shall")


def test_blanks_anchors_and_snake_case_never_atomize():
    # These must tokenize exactly as the plain word split does.
    for text in (f"Price {BLANK} due", "see {#_Ref123_x} here", "a snake_case_name b"):
        assert _tokens(text) == list(dict.fromkeys(t for t in re.split(WORD_SPLIT, text) if t))


def test_adding_defined_term_markers_yields_balanced_hunks():
    original = "The Company shall pay the Investor."
    modified = "The **_Company_** shall pay the **_Investor_**."

    for edit in generate_edits_from_text(original, modified):
        new_text = edit.new_text or ""
        assert new_text.count("**") % 2 == 0, new_text
        assert new_text.count("_") % 2 == 0, new_text


# --- Fix 3(a): tab-stop definitions are not content -------------------------


def test_accept_all_drops_deleted_paragraph_with_tab_stops():
    engine = _engine("Keep me.", "Delete me.")
    p = engine.doc.paragraphs[1]
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1))

    # Track-delete the run and the paragraph mark.
    run_el = p.runs[0]._r
    wrapper = OxmlElement("w:del")
    wrapper.set(qn("w:id"), "901")
    wrapper.set(qn("w:author"), "QA Bot")
    run_el.addprevious(wrapper)
    wrapper.append(run_el)
    for t in run_el.findall(qn("w:t")):
        t.tag = qn("w:delText")

    mark_rpr = OxmlElement("w:rPr")
    mark = OxmlElement("w:del")
    mark.set(qn("w:id"), "902")
    mark.set(qn("w:author"), "QA Bot")
    mark_rpr.append(mark)
    p._p.get_or_add_pPr().append(mark_rpr)

    assert _accepted_paragraphs(engine) == [("Normal", "Keep me.")]


# --- Fix 3(b): positional whole-paragraph replace ---------------------------


def _heading_doc() -> RedlineEngine:
    return _engine("Intro.", ("Costs", "Heading 2"), "Body.")


def _deleted_mark_count(engine: RedlineEngine) -> int:
    return len(engine.doc.element.body.findall(f".//{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}"))


def test_positional_heading_replace_matches_string_path():
    replacement = "1. Expenses"

    # Positional: the internal text diff pins the edit by index.
    positional = _heading_doc()
    original = positional.mapper.full_text
    edits = generate_edits_from_text(original, original.replace("## Costs", replacement))
    assert all(e._match_start_index is not None for e in edits)
    positional.process_batch(edits)

    # String-matched: resolved by _pre_resolve_heuristic_edit.
    matched = _heading_doc()
    matched.process_batch([ModifyText(type="modify", target_text="## Costs", new_text=replacement)])

    assert _deleted_mark_count(positional) == 1
    assert (
        _accepted_paragraphs(positional)
        == _accepted_paragraphs(matched)
        == [
            ("Normal", "Intro."),
            ("List Number", "Expenses"),
            ("Normal", "Body."),
        ]
    )


# --- Fix 4: trailing spaces are content -------------------------------------


def test_markdown_style_keeps_trailing_spaces():
    engine = _engine("x")

    assert engine._parse_markdown_style("## entity.  ") == ("entity.  ", "Heading 2")
    assert engine._parse_markdown_style("- item  ") == ("item  ", "List Paragraph")
    assert engine._parse_markdown_style("1. item  ") == ("item  ", "List Number")
